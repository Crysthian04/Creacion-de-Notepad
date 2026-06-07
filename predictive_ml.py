"""
Análisis Predictivo de Fallos de Máquinas
Combina: SMOTE, GridSearchCV, visualizaciones y exportación a Word
"""

# ── Instalación de dependencias (descomentar si es necesario en Colab) ──────
# !pip install python-docx imbalanced-learn

import io
import sys

import pandas as pd
import matplotlib.pyplot as plt
import seaborn as sns

from imblearn.over_sampling import SMOTE
from sklearn.model_selection import train_test_split, GridSearchCV
from sklearn.tree import DecisionTreeClassifier
from sklearn.ensemble import RandomForestClassifier
from sklearn.metrics import accuracy_score, classification_report, confusion_matrix
from sklearn.preprocessing import LabelEncoder

from docx import Document
from docx.shared import Inches


# ── Utilidades ───────────────────────────────────────────────────────────────

def capture_df_info(df: pd.DataFrame) -> str:
    """Captura la salida de df.info() usando el parámetro buf (seguro ante excepciones)."""
    buffer = io.StringIO()
    df.info(buf=buffer)
    return buffer.getvalue()


def df_to_word_table(doc: Document, df: pd.DataFrame) -> None:
    """Inserta un DataFrame como tabla con formato en el documento Word."""
    table = doc.add_table(rows=1, cols=len(df.columns))
    table.style = "Table Grid"
    for i, col in enumerate(df.columns):
        table.rows[0].cells[i].text = str(col)
    for _, row in df.iterrows():
        cells = table.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = str(round(val, 4) if isinstance(val, float) else val)


def save_figure(path: str) -> None:
    plt.tight_layout()
    plt.savefig(path, dpi=150)
    plt.close()


# ── Carga y exploración de datos ─────────────────────────────────────────────

df = pd.read_csv("Predictive.csv")

print("Primeras filas del DataFrame:")
print(df.head())
print("\nInformación del DataFrame:")
df.info()
print("\nEstadísticas Descriptivas:")
print(df.describe())

# Codificar columnas categóricas
for col in df.select_dtypes(include=["object"]).columns:
    le = LabelEncoder()
    df[col] = le.fit_transform(df[col])

print("\nTipos de datos tras codificación:")
print(df.dtypes)

# ── Matriz de correlación ────────────────────────────────────────────────────

plt.figure(figsize=(12, 6))
sns.heatmap(df.select_dtypes(include=["number"]).corr(), annot=True, cmap="coolwarm", fmt=".2f")
plt.title("Matriz de Correlación entre Variables Numéricas")
save_figure("correlation_matrix.png")
plt.show()

# ── Preparación de datos ─────────────────────────────────────────────────────

target_col = "Machine failure"
extra_targets = ["TWF", "HDF", "PWF", "OSF", "RNF"]

cols_to_drop = [c for c in [target_col] + extra_targets if c in df.columns]
X = df.drop(columns=cols_to_drop)
y = df[target_col]

X_train, X_test, y_train, y_test = train_test_split(X, y, test_size=0.2, random_state=42)

# Balanceo de clases con SMOTE
smote = SMOTE(random_state=42)
X_train_bal, y_train_bal = smote.fit_resample(X_train, y_train)
print(f"\nClases tras SMOTE: {dict(zip(*[list(v) for v in pd.Series(y_train_bal).value_counts().reset_index().values.T]))}")

# ── Árbol de Decisión con GridSearchCV ───────────────────────────────────────

param_grid_tree = {
    "max_depth": [3, 5, 10, None],
    "min_samples_split": [2, 5, 10],
    "min_samples_leaf": [1, 2, 4],
}
grid_tree = GridSearchCV(DecisionTreeClassifier(random_state=42), param_grid_tree, cv=5, scoring="accuracy", n_jobs=-1)
grid_tree.fit(X_train_bal, y_train_bal)

best_tree = grid_tree.best_estimator_
print(f"\nMejores parámetros — Árbol de Decisión: {grid_tree.best_params_}")

y_pred_tree = best_tree.predict(X_test)
acc_tree = accuracy_score(y_test, y_pred_tree)
report_tree = classification_report(y_test, y_pred_tree)
print(f"Exactitud Árbol de Decisión: {acc_tree:.4f}")
print(report_tree)

plt.figure(figsize=(6, 5))
sns.heatmap(confusion_matrix(y_test, y_pred_tree), annot=True, fmt="d", cmap="Blues",
            xticklabels=["No Fallo", "Fallo"], yticklabels=["No Fallo", "Fallo"])
plt.title("Matriz de Confusión — Árbol de Decisión")
plt.xlabel("Predicción")
plt.ylabel("Realidad")
save_figure("conf_matrix_tree.png")
plt.show()

# ── Random Forest con GridSearchCV ───────────────────────────────────────────

param_grid_rf = {
    "n_estimators": [100, 200, 300],
    "max_depth": [5, 10, None],
    "min_samples_split": [2, 5, 10],
}
grid_rf = GridSearchCV(RandomForestClassifier(random_state=42), param_grid_rf, cv=5, scoring="accuracy", n_jobs=-1)
grid_rf.fit(X_train_bal, y_train_bal)

best_rf = grid_rf.best_estimator_
print(f"\nMejores parámetros — Random Forest: {grid_rf.best_params_}")

y_pred_rf = best_rf.predict(X_test)
acc_rf = accuracy_score(y_test, y_pred_rf)
report_rf = classification_report(y_test, y_pred_rf)
print(f"Exactitud Random Forest: {acc_rf:.4f}")
print(report_rf)

plt.figure(figsize=(6, 5))
sns.heatmap(confusion_matrix(y_test, y_pred_rf), annot=True, fmt="d", cmap="Blues",
            xticklabels=["No Fallo", "Fallo"], yticklabels=["No Fallo", "Fallo"])
plt.title("Matriz de Confusión — Random Forest")
plt.xlabel("Predicción")
plt.ylabel("Realidad")
save_figure("conf_matrix_rf.png")
plt.show()

# ── Exportación a Word ───────────────────────────────────────────────────────

doc = Document()
doc.add_heading("Informe de Resultados — Modelos de Machine Learning", 0)

# Exploración de datos
doc.add_heading("1. Primeras filas del DataFrame", level=1)
df_to_word_table(doc, df.head())

doc.add_heading("2. Información del DataFrame", level=1)
doc.add_paragraph(capture_df_info(df))

doc.add_heading("3. Estadísticas Descriptivas", level=1)
df_to_word_table(doc, df.describe().round(4))

doc.add_heading("4. Matriz de Correlación", level=1)
doc.add_picture("correlation_matrix.png", width=Inches(5.5))

# Árbol de Decisión
doc.add_heading("5. Árbol de Decisión", level=1)
doc.add_paragraph(f"Mejores hiperparámetros: {grid_tree.best_params_}")
doc.add_paragraph(f"Exactitud en prueba: {acc_tree:.4f}")
doc.add_paragraph("Reporte de Clasificación:\n" + report_tree)
doc.add_heading("Matriz de Confusión — Árbol de Decisión", level=2)
doc.add_picture("conf_matrix_tree.png", width=Inches(4.5))

# Random Forest
doc.add_heading("6. Random Forest", level=1)
doc.add_paragraph(f"Mejores hiperparámetros: {grid_rf.best_params_}")
doc.add_paragraph(f"Exactitud en prueba: {acc_rf:.4f}")
doc.add_paragraph("Reporte de Clasificación:\n" + report_rf)
doc.add_heading("Matriz de Confusión — Random Forest", level=2)
doc.add_picture("conf_matrix_rf.png", width=Inches(4.5))

# Comparativa final
doc.add_heading("7. Comparativa de Modelos", level=1)
doc.add_paragraph(
    f"{'Modelo':<25} {'Exactitud':>10}\n"
    f"{'─' * 36}\n"
    f"{'Árbol de Decisión':<25} {acc_tree:>10.4f}\n"
    f"{'Random Forest':<25} {acc_rf:>10.4f}\n"
    f"\nModelo recomendado: {'Random Forest' if acc_rf >= acc_tree else 'Árbol de Decisión'}"
)

doc.save("informe_modelos_ml.docx")
print("\nInforme guardado en 'informe_modelos_ml.docx'")
