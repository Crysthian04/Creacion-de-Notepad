"""
Análisis Predictivo — Almacén de Refacciones
Predice: demanda de piezas, riesgo de fallo y quiebre de stock
"""

# ── Instalación de dependencias (descomentar en Colab) ───────────────────────
# !pip install python-docx imbalanced-learn

import io
import numpy as np
import pandas as pd
import matplotlib.pyplot as plt
import seaborn as sns

from imblearn.over_sampling import SMOTE
from sklearn.model_selection import train_test_split, GridSearchCV
from sklearn.tree import DecisionTreeClassifier
from sklearn.ensemble import RandomForestClassifier, RandomForestRegressor
from sklearn.metrics import (
    accuracy_score, classification_report, confusion_matrix,
    mean_absolute_error, mean_squared_error, r2_score,
)
from sklearn.preprocessing import LabelEncoder

from docx import Document
from docx.shared import Inches


# ════════════════════════════════════════════════════════════════════════════
# 1. DATOS DE EJEMPLO
#    Si ya tienes un CSV, reemplaza esta sección con:
#    df = pd.read_csv("tu_archivo.csv")
# ════════════════════════════════════════════════════════════════════════════

np.random.seed(42)
N = 500

codigos = [f"REF-{str(i).zfill(4)}" for i in range(1, 51)]  # 50 tipos de piezas

df = pd.DataFrame({
    "codigo_pieza":       np.random.choice(codigos, N),
    "stock_actual":       np.random.randint(0, 100, N),
    "stock_minimo":       np.random.randint(5, 20, N),
    "salidas_mes":        np.random.poisson(lam=15, size=N),   # demanda mensual
    "dias_sin_movimiento": np.random.randint(0, 90, N),
    "precio_unitario":    np.round(np.random.uniform(50, 5000, N), 2),
    "dias_reposicion":    np.random.randint(1, 30, N),         # lead time proveedor
    "antiguedad_pieza":   np.random.randint(1, 120, N),        # meses en uso
    "categoria":          np.random.choice(["Hidráulica", "Eléctrica", "Mecánica", "Neumática"], N),
})

# Variable objetivo 1 — Quiebre de stock (1 = sí, 0 = no)
df["quiebre_stock"] = ((df["stock_actual"] < df["stock_minimo"]) | (df["stock_actual"] == 0)).astype(int)

# Variable objetivo 2 — Riesgo de fallo de pieza (1 = alto riesgo, 0 = bajo)
df["riesgo_fallo"] = (
    (df["antiguedad_pieza"] > 60) &
    (df["dias_sin_movimiento"] < 10) &
    (df["salidas_mes"] > 18)
).astype(int)

# Variable objetivo 3 — Demanda próximo mes (regresión, valor continuo)
df["demanda_siguiente_mes"] = (
    df["salidas_mes"] * np.random.uniform(0.8, 1.3, N)
).round().astype(int)

print("Dataset generado:")
print(df.head(10))
print(f"\nFilas: {len(df)} | Columnas: {len(df.columns)}")
print(f"\nQuiebre de stock → {df['quiebre_stock'].value_counts().to_dict()}")
print(f"Riesgo de fallo  → {df['riesgo_fallo'].value_counts().to_dict()}")


# ════════════════════════════════════════════════════════════════════════════
# 2. UTILIDADES
# ════════════════════════════════════════════════════════════════════════════

def capture_df_info(df: pd.DataFrame) -> str:
    buffer = io.StringIO()
    df.info(buf=buffer)
    return buffer.getvalue()


def df_to_word_table(doc: Document, df: pd.DataFrame) -> None:
    table = doc.add_table(rows=1, cols=len(df.columns))
    table.style = "Table Grid"
    for i, col in enumerate(df.columns):
        table.rows[0].cells[i].text = str(col)
    for _, row in df.iterrows():
        cells = table.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = str(round(val, 4) if isinstance(val, float) else val)


def save_figure(path: str) -> None:
    plt.tight_layout()
    plt.savefig(path, dpi=150)
    plt.close()


def plot_confusion(y_test, y_pred, title: str, path: str) -> None:
    plt.figure(figsize=(5, 4))
    sns.heatmap(confusion_matrix(y_test, y_pred), annot=True, fmt="d",
                cmap="Blues", xticklabels=["No", "Sí"], yticklabels=["No", "Sí"])
    plt.title(title)
    plt.xlabel("Predicción")
    plt.ylabel("Realidad")
    save_figure(path)


# ════════════════════════════════════════════════════════════════════════════
# 3. PREPROCESAMIENTO
# ════════════════════════════════════════════════════════════════════════════

df_model = df.copy()

# Codificar categoría y código de pieza
for col in df_model.select_dtypes(include=["object"]).columns:
    df_model[col] = LabelEncoder().fit_transform(df_model[col])

# Matriz de correlación
plt.figure(figsize=(13, 7))
sns.heatmap(df_model.select_dtypes(include=["number"]).corr(), annot=True, cmap="coolwarm", fmt=".2f")
plt.title("Matriz de Correlación — Almacén de Refacciones")
save_figure("correlacion_almacen.png")

# Features compartidas para clasificación
FEATURES_CLASIF = [
    "stock_actual", "stock_minimo", "salidas_mes",
    "dias_sin_movimiento", "precio_unitario", "dias_reposicion",
    "antiguedad_pieza", "categoria",
]

PARAM_TREE = {
    "max_depth": [3, 5, 10, None],
    "min_samples_split": [2, 5, 10],
    "min_samples_leaf": [1, 2, 4],
}
PARAM_RF_CLASIF = {
    "n_estimators": [100, 200],
    "max_depth": [5, 10, None],
    "min_samples_split": [2, 5],
}


def clasificacion(nombre: str, X: pd.DataFrame, y: pd.Series):
    """Entrena Árbol de Decisión y Random Forest con SMOTE y GridSearchCV."""
    print(f"\n{'═'*60}")
    print(f"  CLASIFICACIÓN: {nombre}")
    print(f"{'═'*60}")

    X_train, X_test, y_train, y_test = train_test_split(X, y, test_size=0.2, random_state=42)

    smote = SMOTE(random_state=42)
    X_bal, y_bal = smote.fit_resample(X_train, y_train)

    # Árbol de Decisión
    g_tree = GridSearchCV(DecisionTreeClassifier(random_state=42), PARAM_TREE, cv=5, scoring="accuracy", n_jobs=-1)
    g_tree.fit(X_bal, y_bal)
    best_tree = g_tree.best_estimator_
    y_pred_tree = best_tree.predict(X_test)
    acc_tree = accuracy_score(y_test, y_pred_tree)
    rep_tree = classification_report(y_test, y_pred_tree)
    print(f"[Árbol] Parámetros: {g_tree.best_params_} | Exactitud: {acc_tree:.4f}")

    # Random Forest
    g_rf = GridSearchCV(RandomForestClassifier(random_state=42), PARAM_RF_CLASIF, cv=5, scoring="accuracy", n_jobs=-1)
    g_rf.fit(X_bal, y_bal)
    best_rf = g_rf.best_estimator_
    y_pred_rf = best_rf.predict(X_test)
    acc_rf = accuracy_score(y_test, y_pred_rf)
    rep_rf = classification_report(y_test, y_pred_rf)
    print(f"[RF]    Parámetros: {g_rf.best_params_}  | Exactitud: {acc_rf:.4f}")

    tag = nombre.lower().replace(" ", "_")
    plot_confusion(y_test, y_pred_tree, f"Árbol — {nombre}", f"cm_tree_{tag}.png")
    plot_confusion(y_test, y_pred_rf,   f"Random Forest — {nombre}", f"cm_rf_{tag}.png")

    return {
        "nombre": nombre, "tag": tag,
        "tree_params": g_tree.best_params_, "tree_acc": acc_tree, "tree_rep": rep_tree,
        "rf_params": g_rf.best_params_,     "rf_acc": acc_rf,     "rf_rep": rep_rf,
        "y_test": y_test, "y_pred_tree": y_pred_tree, "y_pred_rf": y_pred_rf,
    }


# ════════════════════════════════════════════════════════════════════════════
# 4. ANÁLISIS 1 — QUIEBRE DE STOCK (clasificación)
# ════════════════════════════════════════════════════════════════════════════

X_qs = df_model[FEATURES_CLASIF]
y_qs = df_model["quiebre_stock"]
res_qs = clasificacion("Quiebre de Stock", X_qs, y_qs)

# ════════════════════════════════════════════════════════════════════════════
# 5. ANÁLISIS 2 — RIESGO DE FALLO (clasificación)
# ════════════════════════════════════════════════════════════════════════════

X_rf = df_model[FEATURES_CLASIF]
y_rf_fallo = df_model["riesgo_fallo"]
res_fallo = clasificacion("Riesgo de Fallo de Pieza", X_rf, y_rf_fallo)

# ════════════════════════════════════════════════════════════════════════════
# 6. ANÁLISIS 3 — DEMANDA SIGUIENTE MES (regresión)
# ════════════════════════════════════════════════════════════════════════════

print(f"\n{'═'*60}")
print("  REGRESIÓN: Demanda Siguiente Mes")
print(f"{'═'*60}")

FEATURES_REG = ["stock_actual", "salidas_mes", "dias_sin_movimiento",
                "precio_unitario", "dias_reposicion", "antiguedad_pieza", "categoria"]

X_dem = df_model[FEATURES_REG]
y_dem = df_model["demanda_siguiente_mes"]

X_tr_d, X_te_d, y_tr_d, y_te_d = train_test_split(X_dem, y_dem, test_size=0.2, random_state=42)

rf_reg = RandomForestRegressor(n_estimators=200, random_state=42, n_jobs=-1)
rf_reg.fit(X_tr_d, y_tr_d)
y_pred_dem = rf_reg.predict(X_te_d)

mae  = mean_absolute_error(y_te_d, y_pred_dem)
rmse = np.sqrt(mean_squared_error(y_te_d, y_pred_dem))
r2   = r2_score(y_te_d, y_pred_dem)
print(f"MAE: {mae:.2f} | RMSE: {rmse:.2f} | R²: {r2:.4f}")

# Gráfico real vs predicho
plt.figure(figsize=(7, 4))
plt.scatter(y_te_d, y_pred_dem, alpha=0.5, color="steelblue")
plt.plot([y_te_d.min(), y_te_d.max()], [y_te_d.min(), y_te_d.max()], "r--")
plt.xlabel("Demanda Real")
plt.ylabel("Demanda Predicha")
plt.title("Demanda Siguiente Mes — Real vs Predicha")
save_figure("demanda_real_vs_pred.png")

# Importancia de variables
feat_imp = pd.Series(rf_reg.feature_importances_, index=FEATURES_REG).sort_values(ascending=False)
plt.figure(figsize=(7, 4))
feat_imp.plot(kind="bar", color="steelblue")
plt.title("Importancia de Variables — Demanda")
plt.ylabel("Importancia")
save_figure("importancia_demanda.png")


# ════════════════════════════════════════════════════════════════════════════
# 7. EXPORTACIÓN A WORD
# ════════════════════════════════════════════════════════════════════════════

doc = Document()
doc.add_heading("Informe Predictivo — Almacén de Refacciones", 0)

# ── Exploración ──────────────────────────────────────────────────────────────
doc.add_heading("1. Vista del Dataset", level=1)
df_to_word_table(doc, df.head(10))

doc.add_heading("2. Información del DataFrame", level=1)
doc.add_paragraph(capture_df_info(df))

doc.add_heading("3. Estadísticas Descriptivas", level=1)
df_to_word_table(doc, df.describe().round(2))

doc.add_heading("4. Matriz de Correlación", level=1)
doc.add_picture("correlacion_almacen.png", width=Inches(6))


def agregar_seccion_clasificacion(doc, res: dict, num: int) -> None:
    nombre = res["nombre"]
    tag = res["tag"]

    doc.add_heading(f"{num}. {nombre}", level=1)

    doc.add_heading("Árbol de Decisión", level=2)
    doc.add_paragraph(f"Mejores hiperparámetros: {res['tree_params']}")
    doc.add_paragraph(f"Exactitud: {res['tree_acc']:.4f}")
    doc.add_paragraph("Reporte de Clasificación:\n" + res["tree_rep"])
    doc.add_picture(f"cm_tree_{tag}.png", width=Inches(4))

    doc.add_heading("Random Forest", level=2)
    doc.add_paragraph(f"Mejores hiperparámetros: {res['rf_params']}")
    doc.add_paragraph(f"Exactitud: {res['rf_acc']:.4f}")
    doc.add_paragraph("Reporte de Clasificación:\n" + res["rf_rep"])
    doc.add_picture(f"cm_rf_{tag}.png", width=Inches(4))

    mejor = "Random Forest" if res["rf_acc"] >= res["tree_acc"] else "Árbol de Decisión"
    doc.add_paragraph(f"Modelo recomendado para {nombre}: {mejor}")


agregar_seccion_clasificacion(doc, res_qs, 5)
agregar_seccion_clasificacion(doc, res_fallo, 6)

# ── Demanda ──────────────────────────────────────────────────────────────────
doc.add_heading("7. Predicción de Demanda Siguiente Mes", level=1)
doc.add_paragraph(f"Modelo: Random Forest Regressor (n_estimators=200)")
doc.add_paragraph(
    f"MAE  (Error Absoluto Medio):  {mae:.2f} unidades\n"
    f"RMSE (Raíz del Error Cuad.):  {rmse:.2f} unidades\n"
    f"R²   (Coef. de determinación): {r2:.4f}"
)
doc.add_picture("demanda_real_vs_pred.png", width=Inches(5.5))
doc.add_heading("Importancia de Variables — Demanda", level=2)
doc.add_picture("importancia_demanda.png", width=Inches(5.5))

# ── Comparativa ──────────────────────────────────────────────────────────────
doc.add_heading("8. Resumen Comparativo", level=1)
resumen = pd.DataFrame([
    {"Análisis": "Quiebre de Stock",      "Modelo": "Árbol de Decisión", "Exactitud": f"{res_qs['tree_acc']:.4f}"},
    {"Análisis": "Quiebre de Stock",      "Modelo": "Random Forest",     "Exactitud": f"{res_qs['rf_acc']:.4f}"},
    {"Análisis": "Riesgo de Fallo",       "Modelo": "Árbol de Decisión", "Exactitud": f"{res_fallo['tree_acc']:.4f}"},
    {"Análisis": "Riesgo de Fallo",       "Modelo": "Random Forest",     "Exactitud": f"{res_fallo['rf_acc']:.4f}"},
    {"Análisis": "Demanda Siguiente Mes", "Modelo": "Random Forest Reg.", "Exactitud": f"R²={r2:.4f}"},
])
df_to_word_table(doc, resumen)

doc.save("informe_almacen_refacciones.docx")
print("\nInforme guardado en 'informe_almacen_refacciones.docx'")
