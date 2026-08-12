"""`Parametros`: every configurable constant of the workbook.

Two regions on one sheet:

  left  -- parameter blocks, one row per key, each value cell reachable through
           a `PARAM_<clave>` defined name;
  right -- catalogue columns, each reachable through a `CAT_<catalogo>` defined
           name and used as the source of the list validations of every other
           sheet.

No other sheet may contain a business constant. When a formula elsewhere needs a
threshold, a limit or a code, it reads it from here by name.

Catalogue ranges are sized exactly to their seeded content. An approver who adds
a code by hand must extend the range; phase 9 adds a button that does it.
"""

from __future__ import annotations

from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from src.common import names, styles
from src.common.layout import Column, SheetLayout
from src.sheets.base import HIDDEN, BuildContext, SheetSpec

SPEC = SheetSpec(key="parametros", title="Parametros", visibility=HIDDEN, role="APROBADOR")

# Spanish captions of the parameter blocks, in display order. A block present in
# the seed but absent here would be silently dropped, so the builder checks.
BLOCK_TITLES = {
    "IDENTIFICACION": "Identificación y correlativos",
    "CRITICIDAD_PESOS": "Pesos y escalas de la matriz de criticidad",
    "CRITICIDAD_UMBRALES": "Umbrales de criticidad",
    "ADAPTADOR": "Adaptador CMMS activo",
    "LIMITES_TEXTO": "Límites de caracteres por adaptador",
    "SAP": "Parámetros de SAP PM",
    "MAXIMO": "Parámetros de IBM Maximo",
    "EMPAQUETADO": "Parámetros de empaquetado",
    "SEGURIDAD": "Roles y contraseñas",
    "VALIDACION": "Reglas de validación",
}

# Region keys other builders use to find the catalogue ranges.
REGION_CATALOG_FIRST_COLUMN = "parametros.catalogos.primera_columna"

_FIRST_CATALOG_COLUMN = 7  # column G, clear of the parameter block
_NUMERIC_TYPES = {"ENTERO", "DECIMAL"}


class ParametrosSheetBuilder:
    """Builds the parameter and catalogue sheet."""

    spec = SPEC

    def build(self, worksheet: Worksheet, context: BuildContext) -> None:
        self._check_blocks(context)
        self._build_catalogs(worksheet, context)
        self._build_parameters(worksheet, context)
        self._build_task_types(worksheet, context)

    # -- parameters ------------------------------------------------------

    def _check_blocks(self, context: BuildContext) -> None:
        unknown = set(context.seed.parametros_por_bloque) - set(BLOCK_TITLES)
        if unknown:
            raise ValueError(
                "parametros.csv declares block(s) with no caption in BLOCK_TITLES: "
                + ", ".join(sorted(unknown))
            )

    def _build_parameters(self, worksheet: Worksheet, context: BuildContext) -> None:
        layout = SheetLayout(worksheet, first_column=2)
        layout.title(
            "Parámetros del libro",
            "Toda constante de negocio vive aquí. Solo el rol Aprobador debería editarla.",
        )

        for block, title in BLOCK_TITLES.items():
            rows = context.seed.parametros_por_bloque.get(block)
            if not rows:
                continue
            layout.section(title, width=3)
            first_row = layout.row
            for index, row in enumerate(rows):
                self._write_parameter(worksheet, layout, first_row + index, row, context)
            layout.row = first_row + len(rows) + 1

    def _write_parameter(
        self,
        worksheet: Worksheet,
        layout: SheetLayout,
        row: int,
        parameter: dict[str, str],
        context: BuildContext,
    ) -> None:
        label_col = layout.first_column
        value_col = label_col + 1
        note_col = label_col + 2

        label_cell = worksheet.cell(row=row, column=label_col, value=parameter["etiqueta"])
        label_cell.style = styles.LABEL

        value_cell = worksheet.cell(row=row, column=value_col)
        value_cell.style = self._value_style(parameter["tipo"])
        value = self._typed_value(parameter)
        if value is not None:
            value_cell.value = value

        note_cell = worksheet.cell(row=row, column=note_col)
        note_cell.style = styles.NOTE
        note_cell.value = parameter["notas"] or parameter["clave"]

        names.define(
            context.workbook,
            names.param_name(parameter["clave"]),
            names.cell_ref(self.spec.title, value_col, row),
        )

        catalogue = parameter["catalogo"]
        if catalogue:
            from src.common.layout import add_list_validation

            add_list_validation(
                worksheet, names.catalog_name(catalogue), row, row, value_col, allow_blank=False
            )

    @staticmethod
    def _value_style(tipo: str) -> str:
        if tipo == "ENTERO":
            return styles.INPUT_NUMBER
        if tipo == "DECIMAL":
            return styles.INPUT_DECIMAL
        if tipo == "HASH":
            return styles.CODE
        return styles.INPUT_TEXT

    @staticmethod
    def _typed_value(parameter: dict[str, str]) -> object | None:
        raw = parameter["valor"]
        if raw == "":
            return None
        if parameter["tipo"] not in _NUMERIC_TYPES:
            return raw
        try:
            return int(raw) if parameter["tipo"] == "ENTERO" else float(raw)
        except ValueError as exc:
            raise ValueError(
                f"Parameter '{parameter['clave']}' declares type {parameter['tipo']} "
                f"but its value '{raw}' is not numeric."
            ) from exc

    # -- catalogues ------------------------------------------------------

    def _build_catalogs(self, worksheet: Worksheet, context: BuildContext) -> None:
        """Write one two-column block per catalogue: code and label.

        `CAT_<name>` covers the code column only, because that is what a data
        validation list must yield; the label column is there for the analyst
        reading the sheet.
        """
        column = _FIRST_CATALOG_COLUMN
        header_row = 5
        worksheet.cell(row=header_row - 2, column=column, value="Catálogos").style = styles.TITLE

        for catalogue, entries in sorted(context.seed.catalogos_por_nombre.items()):
            code_cell = worksheet.cell(row=header_row, column=column, value=catalogue)
            code_cell.style = styles.HEADER
            label_cell = worksheet.cell(row=header_row, column=column + 1, value="Etiqueta")
            label_cell.style = styles.HEADER
            worksheet.column_dimensions[get_column_letter(column)].width = 18
            worksheet.column_dimensions[get_column_letter(column + 1)].width = 34

            for index, entry in enumerate(entries):
                row = header_row + 1 + index
                worksheet.cell(row=row, column=column, value=entry["codigo"]).style = styles.CODE
                worksheet.cell(
                    row=row, column=column + 1, value=entry["etiqueta"]
                ).style = styles.READONLY

            names.define(
                context.workbook,
                names.catalog_name(catalogue),
                names.range_ref(
                    self.spec.title, column, header_row + 1, column, header_row + len(entries)
                ),
            )
            column += 2

        context.publish(REGION_CATALOG_FIRST_COLUMN, str(_FIRST_CATALOG_COLUMN))

    # -- task type attributes -------------------------------------------

    def _build_task_types(self, worksheet: Worksheet, context: BuildContext) -> None:
        """Attributes of each task type: MSG-3 category and admissible interval.

        The decision engine and validator rule 8 read these; they are attributes
        of the task type rather than free parameters, so they get their own table.
        """
        columns = [
            Column("codigo", "Código", 12, styles.CODE),
            Column("etiqueta", "Tipo de tarea", 40, styles.READONLY),
            Column("categoria_msg3", "Categoría MSG-3", 18, styles.CODE),
            Column("codigo_actividad", "Código de actividad", 18, styles.CODE),
            Column("oficio_sugerido", "Oficio sugerido", 16, styles.CODE),
            Column("requiere_intervalo", "Requiere intervalo", 14, styles.CODE),
            Column("requiere_justificacion", "Requiere justificación", 16, styles.CODE),
            Column("frecuencia_min_dias", "Frecuencia mínima (días)", 16, styles.INPUT_NUMBER),
            Column("frecuencia_max_dias", "Frecuencia máxima (días)", 16, styles.INPUT_NUMBER),
            Column("activo", "Activo", 10, styles.CODE),
        ]

        rows = [
            {
                **row,
                "frecuencia_min_dias": int(row["frecuencia_min_dias"]),
                "frecuencia_max_dias": int(row["frecuencia_max_dias"]),
            }
            for row in context.seed.tipos_tarea
        ]

        # Placed below the catalogue block, on its own row band.
        start_row = 7 + max(len(entries) for entries in context.seed.catalogos_por_nombre.values())
        layout = SheetLayout(worksheet, row=start_row, first_column=_FIRST_CATALOG_COLUMN)
        layout.section("Atributos de los tipos de tarea", width=len(columns))
        region = layout.table(columns, rows)

        names.define(
            context.workbook,
            names.range_name("TiposTarea"),
            names.range_ref(
                self.spec.title,
                region.first_column,
                region.first_data_row,
                region.last_column,
                region.last_data_row,
            ),
        )
