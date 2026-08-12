"""Sheet builder contract and registry.

Every worksheet of the template is produced by one builder object. The registry
in `src/sheets/__init__.py` defines the physical tab order: builders are applied
in list order, so reordering the list reorders the tabs.

A builder receives a `BuildContext`, which is everything it may legitimately
read: the workbook (to register defined names), the build configuration and the
seed data. A builder never reads a file or a global.
"""

from __future__ import annotations

from dataclasses import dataclass, field
from typing import Protocol

from openpyxl.workbook.workbook import Workbook
from openpyxl.worksheet.worksheet import Worksheet

from src.common.config import BuildConfig
from src.seedloader import SeedData

# Excel sheet visibility states, mirrored here so builders never import the
# raw string literals.
VISIBLE = "visible"
HIDDEN = "hidden"
VERY_HIDDEN = "veryHidden"


@dataclass(frozen=True)
class SheetSpec:
    """Static description of a sheet, independent of its contents.

    `title` is the Spanish tab caption seen by the analyst; `key` is the stable
    English identifier used by code and tests, which never changes even if the
    caption is retitled or translated.
    """

    key: str
    title: str
    visibility: str = VISIBLE
    role: str = "LECTOR"

    def __post_init__(self) -> None:
        if self.visibility not in (VISIBLE, HIDDEN, VERY_HIDDEN):
            raise ValueError(f"Unknown visibility '{self.visibility}' for sheet '{self.key}'.")
        # Excel's own limit; catching it here beats a corrupt file later.
        if not self.title or len(self.title) > 31:
            raise ValueError(f"Sheet title must be 1-31 characters: '{self.title}'.")


@dataclass
class BuildContext:
    """Everything a builder may read, plus the regions builders publish.

    `regions` lets a later builder reference a range an earlier one produced --
    for example the task sheet pointing at the catalogue ranges of `Parametros`
    -- without either of them knowing the other's layout.
    """

    workbook: Workbook
    config: BuildConfig
    seed: SeedData
    regions: dict[str, str] = field(default_factory=dict)

    def publish(self, key: str, reference: str) -> None:
        if key in self.regions:
            raise KeyError(f"Region '{key}' was already published.")
        self.regions[key] = reference

    def label(self, clave: str, default: str | None = None) -> str:
        """Spanish label for a dictionary key.

        Sheets take their captions from the dictionary so that translating the
        workbook is a matter of filling in the `en` column.
        """
        for row in self.seed.diccionario:
            if row["clave"] == clave:
                return row["es"]
        if default is not None:
            return default
        raise KeyError(f"Dictionary key '{clave}' is not defined in the seed.")


class SheetBuilder(Protocol):
    """Anything that can materialise one sheet into a workbook."""

    spec: SheetSpec

    def build(self, worksheet: Worksheet, context: BuildContext) -> None:
        """Populate `worksheet`. The sheet already exists and is correctly named."""


def create_sheet(workbook: Workbook, spec: SheetSpec) -> Worksheet:
    """Append a sheet for `spec` and apply its visibility."""
    worksheet = workbook.create_sheet(title=spec.title)
    worksheet.sheet_state = spec.visibility
    return worksheet


def build_all(
    workbook: Workbook, builders: list[SheetBuilder], context: BuildContext
) -> list[SheetSpec]:
    """Apply every builder in order and return the specs that were built."""
    built: list[SheetSpec] = []
    for builder in builders:
        worksheet = create_sheet(workbook, builder.spec)
        builder.build(worksheet, context)
        built.append(builder.spec)
    return built
