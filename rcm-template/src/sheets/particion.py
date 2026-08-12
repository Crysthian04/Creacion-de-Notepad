"""`Particion`: the three level breakdown of the machine.

assembly (1) -> sub-assembly (2) -> maintainable item (3).

The analyst types only the level, the description and the ISO 14224 code. The
item number, the concatenated caption and its length are formulas, and the
hierarchical numbering (`1`, `1.2`, `1.2.04`) is carried down the sheet by three
hidden counter columns -- one per level -- so that no formula nests more than
three levels deep and none of them is volatile.
"""

from __future__ import annotations

from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from src.common import header, names, styles
from src.common.layout import Column, SheetLayout, add_custom_validation, add_range_validation
from src.sheets.base import VISIBLE, BuildContext, SheetSpec

SPEC = SheetSpec(key="particion", title="Particion", visibility=VISIBLE, role="ANALISTA")

# Sized for a single machine: a 300 row breakdown is already a very detailed one.
DATA_ROWS = 300

RANGE_LOOKUP = "Particion"
RANGE_IDS = "ParticionIds"
RANGE_LEVEL3_IDS = "ParticionItems"

COLUMNS = [
    Column("id_item", "ID de ítem", 14, styles.READONLY),
    Column("nivel", "Nivel", 8, styles.INPUT_NUMBER),
    Column("n1", "n1", 6, styles.READONLY),
    Column("n2", "n2", 6, styles.READONLY),
    Column("n3", "n3", 6, styles.READONLY),
    Column("item", "Ítem", 14, styles.READONLY),
    Column("descripcion", "Descripción", 46, styles.INPUT_TEXT),
    Column("codigo_taxonomia", "Código ISO 14224", 20, styles.INPUT_TEXT),
    Column("concatenado", "Concatenado", 60, styles.READONLY),
    Column("longitud", "Longitud", 10, styles.READONLY),
    Column("notas", "Notas", 40, styles.INPUT_TEXT),
]

# Counter columns exist for the numbering only; hiding them keeps the sheet readable.
HIDDEN_KEYS = ("n1", "n2", "n3")


class ParticionSheetBuilder:
    """Builds the machine breakdown sheet."""

    spec = SPEC

    def build(self, worksheet: Worksheet, context: BuildContext) -> None:
        layout = SheetLayout(worksheet, first_column=2)
        layout.title(
            "Partición del activo",
            "Tres niveles: ensamblaje → sub-ensamblaje → componente mantenible. "
            "La numeración es automática.",
        )
        header.build_echo(worksheet, layout)
        layout.note(
            "Escriba el nivel (1, 2 o 3) y la descripción. El ítem, el concatenado y "
            "la longitud se calculan solos."
        )
        layout.blank()

        region = layout.table(COLUMNS, [], blank_rows=DATA_ROWS, autofilter=True, freeze=True)
        column_of = {
            column.key: region.first_column + offset for offset, column in enumerate(COLUMNS)
        }

        self._write_formulas(worksheet, region, column_of)
        self._hide_counters(worksheet, column_of)
        self._validations(worksheet, region, column_of)
        self._define_names(context, region, column_of)

        worksheet.column_dimensions["A"].width = 3

    # -- formulas --------------------------------------------------------

    def _write_formulas(self, worksheet: Worksheet, region, column_of: dict[str, int]) -> None:
        letter = {key: get_column_letter(column) for key, column in column_of.items()}

        for row in range(region.first_data_row, region.last_data_row + 1):
            previous = row - 1
            level = f"${letter['nivel']}{row}"

            # Counters carry the current number of each level down the sheet.
            worksheet.cell(
                row=row, column=column_of["n1"]
            ).value = f"=IF({level}=1,{letter['n1']}{previous}+1,{letter['n1']}{previous})"
            worksheet.cell(row=row, column=column_of["n2"]).value = (
                f"=IF({level}=1,0,IF({level}=2,{letter['n2']}{previous}+1,"
                f"{letter['n2']}{previous}))"
            )
            worksheet.cell(
                row=row, column=column_of["n3"]
            ).value = f"=IF({level}<=2,0,{letter['n3']}{previous}+1)"

            # Built by concatenation rather than by branching on the level, so
            # the formula stays within three levels of nesting.
            n1 = f'TEXT({letter["n1"]}{row},"0")'
            n2 = f'TEXT({letter["n2"]}{row},"0")'
            n3 = f'TEXT({letter["n3"]}{row},"00")'
            worksheet.cell(row=row, column=column_of["item"]).value = (
                f'=IF({level}="","",{n1}&IF({level}>=2,"."&{n2},"")&IF({level}=3,"."&{n3},""))'
            )

            item = f"{letter['item']}{row}"
            description = f"{letter['descripcion']}{row}"
            worksheet.cell(row=row, column=column_of["id_item"]).value = (
                f'=IF({item}="","",PARAM_PREFIJO_ITEM&TEXT(ROW()-{region.header_row},"0000"))'
            )
            worksheet.cell(
                row=row, column=column_of["concatenado"]
            ).value = f'=IF({item}="","",{item}&" - "&{description})'
            worksheet.cell(
                row=row, column=column_of["longitud"]
            ).value = f'=IF({letter["concatenado"]}{row}="","",LEN({letter["concatenado"]}{row}))'

        # Seed the counters on the header row so the first data row has a base.
        for key in HIDDEN_KEYS:
            worksheet.cell(row=region.header_row, column=column_of[key]).value = 0

    @staticmethod
    def _hide_counters(worksheet: Worksheet, column_of: dict[str, int]) -> None:
        for key in HIDDEN_KEYS:
            worksheet.column_dimensions[get_column_letter(column_of[key])].hidden = True

    # -- validations -----------------------------------------------------

    @staticmethod
    def _validations(worksheet: Worksheet, region, column_of: dict[str, int]) -> None:
        add_range_validation(
            worksheet,
            region.first_data_row,
            region.last_data_row,
            column_of["nivel"],
            minimum=1,
            maximum=3,
            error_message="El nivel debe ser 1 (ensamblaje), 2 (sub-ensamblaje) o 3 (componente).",
        )
        taxonomy_cell = f"{get_column_letter(column_of['codigo_taxonomia'])}"
        add_custom_validation(
            worksheet,
            f"COUNTIF(LIB_TAXONOMIACLAVE,{taxonomy_cell}{region.first_data_row})>0",
            region.first_data_row,
            region.last_data_row,
            column_of["codigo_taxonomia"],
            error_title="Código de taxonomía desconocido",
            error_message=("El código debe existir en la taxonomía ISO 14224 cargada en el libro."),
        )

    # -- names -----------------------------------------------------------

    def _define_names(self, context: BuildContext, region, column_of: dict[str, int]) -> None:
        names.define(
            context.workbook,
            names.range_name(RANGE_LOOKUP),
            names.range_ref(
                self.spec.title,
                region.first_column,
                region.first_data_row,
                region.last_column,
                region.last_data_row,
            ),
        )
        names.define(
            context.workbook,
            names.range_name(RANGE_IDS),
            names.range_ref(
                self.spec.title,
                column_of["id_item"],
                region.first_data_row,
                column_of["id_item"],
                region.last_data_row,
            ),
        )
        names.define(
            context.workbook,
            names.range_name("ParticionNivel"),
            names.range_ref(
                self.spec.title,
                column_of["nivel"],
                region.first_data_row,
                column_of["nivel"],
                region.last_data_row,
            ),
        )
        names.define(
            context.workbook,
            names.range_name("ParticionTaxonomia"),
            names.range_ref(
                self.spec.title,
                column_of["codigo_taxonomia"],
                region.first_data_row,
                column_of["codigo_taxonomia"],
                region.last_data_row,
            ),
        )
        # Offset of the description column inside RNG_Particion, for VLOOKUP.
        context.publish(
            "particion.columna_descripcion",
            str(column_of["descripcion"] - region.first_column + 1),
        )
        context.publish("particion.columna_item", str(column_of["item"] - region.first_column + 1))
