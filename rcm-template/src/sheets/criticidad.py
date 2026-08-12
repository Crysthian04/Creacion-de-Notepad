"""`Criticidad del Activo`: the weighted probability x consequence matrix.

Four consequence categories -- EH&S, Calidad, Servicio, Costo -- each rated on
the configured scale and weighted. The sheet holds no number of its own: scale
bounds, weights and class thresholds all resolve to `Parametros` by defined name,
which is what makes the matrix re-tunable without touching a formula.

Score model, kept deliberately flat so no formula nests more than three levels:

    consecuencia ponderada = SUM(consecuencia_i x peso_i)
    puntaje                = probabilidad x consecuencia ponderada
    clase                  = first threshold the score reaches
"""

from __future__ import annotations

from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from src.common import header, names, styles
from src.common.layout import Column, SheetLayout, add_range_validation
from src.sheets.base import VISIBLE, BuildContext, SheetSpec

SPEC = SheetSpec(
    key="criticidad", title="Criticidad del Activo", visibility=VISIBLE, role="ANALISTA"
)

# Category key, caption, and the parameter holding its weight.
CATEGORIES = (
    ("ehs", "EH&S (seguridad, salud y ambiente)", "peso_ehs"),
    ("calidad", "Calidad del producto", "peso_calidad"),
    ("servicio", "Servicio / continuidad operativa", "peso_servicio"),
    ("costo", "Costo directo de la falla", "peso_costo"),
)

RANGE_SCORE = "CriticidadPuntaje"
RANGE_CLASS = "CriticidadClase"
RANGE_PROBABILITY = "CriticidadProbabilidad"


class CriticidadSheetBuilder:
    """Builds the asset criticality matrix."""

    spec = SPEC

    def build(self, worksheet: Worksheet, context: BuildContext) -> None:
        layout = SheetLayout(worksheet, first_column=2)
        layout.title(
            "Criticidad del activo",
            "Pesos, escalas y umbrales se configuran en Parametros; esta hoja no "
            "contiene ningún número fijo.",
        )
        header.build_echo(worksheet, layout)

        probability_row = self._probability(worksheet, layout, context)
        region = self._matrix(worksheet, layout, context)
        self._result(worksheet, layout, context, region, probability_row)

        worksheet.column_dimensions["A"].width = 3

    # -- inputs ----------------------------------------------------------

    def _probability(self, worksheet: Worksheet, layout: SheetLayout, context: BuildContext) -> int:
        layout.section("Probabilidad de la falla", width=3)
        row = layout.row
        value_col = layout.first_column + 1

        layout.key_values(
            [
                (
                    "Probabilidad",
                    None,
                    "Escala configurada en Parametros (escala_min a escala_max).",
                )
            ]
        )
        add_range_validation(
            worksheet,
            row,
            row,
            value_col,
            minimum="PARAM_ESCALA_MIN",
            maximum="PARAM_ESCALA_MAX",
            error_message="La probabilidad debe estar dentro de la escala configurada.",
        )
        names.define(
            context.workbook,
            names.range_name(RANGE_PROBABILITY),
            names.cell_ref(self.spec.title, value_col, row),
        )
        return row

    def _matrix(self, worksheet: Worksheet, layout: SheetLayout, context: BuildContext):
        layout.section("Consecuencia por categoría", width=4)
        columns = [
            Column("categoria", "Categoría", 44, styles.READONLY),
            Column("consecuencia", "Consecuencia", 16, styles.INPUT_NUMBER),
            Column("peso", "Peso", 12, styles.READONLY_NUMBER),
            Column("ponderado", "Consecuencia ponderada", 22, styles.READONLY_NUMBER),
        ]
        rows = [{"categoria": caption} for _, caption, _ in CATEGORIES]
        region = layout.table(columns, rows)

        consequence_col = region.first_column + 1
        weight_col = region.first_column + 2
        weighted_col = region.first_column + 3

        for index, (_, _, weight_param) in enumerate(CATEGORIES):
            row = region.first_data_row + index
            worksheet.cell(row=row, column=weight_col).value = f"={names.param_name(weight_param)}"
            consequence = f"{get_column_letter(consequence_col)}{row}"
            weight = f"{get_column_letter(weight_col)}{row}"
            worksheet.cell(row=row, column=weighted_col).value = f"={consequence}*{weight}"

        add_range_validation(
            worksheet,
            region.first_data_row,
            region.last_data_row,
            consequence_col,
            minimum="PARAM_ESCALA_MIN",
            maximum="PARAM_ESCALA_MAX",
            error_message="La consecuencia debe estar dentro de la escala configurada.",
        )
        return region

    # -- result ----------------------------------------------------------

    def _result(
        self,
        worksheet: Worksheet,
        layout: SheetLayout,
        context: BuildContext,
        region,
        probability_row: int,
    ) -> None:
        weighted_col = get_column_letter(region.first_column + 3)
        weighted_sum = (
            f"SUM({weighted_col}{region.first_data_row}:{weighted_col}{region.last_data_row})"
        )
        probability = f"{get_column_letter(layout.first_column + 1)}{probability_row}"

        layout.section("Resultado", width=3)
        first_row = layout.row
        value_col = layout.first_column + 1

        layout.key_values(
            [
                ("Consecuencia ponderada", None, "Suma de las cuatro categorías."),
                ("Puntaje de criticidad", None, "Probabilidad x consecuencia ponderada."),
                ("Clase sin filtrar", None, "Auxiliar: clase antes de descartar el caso vacío."),
                ("Clase de criticidad", None, "Se hereda a cada tarea del plan."),
            ],
            value_style=styles.READONLY,
        )

        weighted_row = first_row
        score_row = first_row + 1
        raw_class_row = first_row + 2
        class_row = first_row + 3
        score_ref = f"{get_column_letter(value_col)}{score_row}"

        worksheet.cell(row=weighted_row, column=value_col).value = f"={weighted_sum}"
        worksheet.cell(
            row=score_row, column=value_col
        ).value = f"={probability}*{get_column_letter(value_col)}{weighted_row}"
        # The empty-input guard lives in its own cell: folding it into the
        # threshold chain would make a fourth nesting level.
        worksheet.cell(row=raw_class_row, column=value_col).value = (
            f'=IF({score_ref}>=PARAM_UMBRAL_CLASE_A,"A",'
            f'IF({score_ref}>=PARAM_UMBRAL_CLASE_B,"B",'
            f'IF({score_ref}>=PARAM_UMBRAL_CLASE_C,"C","D")))'
        )
        worksheet.cell(
            row=class_row, column=value_col
        ).value = f'=IF({score_ref}=0,"",{get_column_letter(value_col)}{raw_class_row})'
        worksheet.row_dimensions[raw_class_row].hidden = True

        names.define(
            context.workbook,
            names.range_name(RANGE_SCORE),
            names.cell_ref(self.spec.title, value_col, score_row),
        )
        names.define(
            context.workbook,
            names.range_name(RANGE_CLASS),
            names.cell_ref(self.spec.title, value_col, class_row),
        )

        self._colour_class(worksheet, value_col, class_row)

    @staticmethod
    def _colour_class(worksheet: Worksheet, column: int, row: int) -> None:
        """Colour the class cell by its own value, using the shared palette."""
        cell = f"{get_column_letter(column)}{row}"
        for code, colour in styles.CRITICALITY_FILLS.items():
            worksheet.conditional_formatting.add(
                cell,
                CellIsRule(
                    operator="equal",
                    formula=[f'"{code}"'],
                    fill=PatternFill("solid", bgColor=colour),
                    font=Font(bold=True),
                ),
            )
