"""`AMFE`: the functional analysis chain.

función → falla funcional → modo de falla → causa → efecto, one `FM-nnnn` record
per failure mode, plus the four RCM2 consequence questions that feed the decision
tree.

The consequence category is derived on the sheet rather than by macro, so the
workbook still classifies correctly when it is opened without macros. Deriving it
in one formula would need four nesting levels, which the project forbids, so three
hidden helper columns split the logic:

    eje       -> which consequence dominates: SEG / AMB / OPE / NOOPE
    evidencia -> EVIDENTE or OCULTA
    oculta    -> the two-way split a hidden failure collapses to
    categoria -> combines them, and is the only one the analyst sees
"""

from __future__ import annotations

from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from src.common import header, names, styles
from src.common.layout import Column, SheetLayout, add_custom_validation, add_list_validation
from src.sheets.base import VISIBLE, BuildContext, SheetSpec

SPEC = SheetSpec(key="amfe", title="AMFE", visibility=VISIBLE, role="ANALISTA")

DATA_ROWS = 500

RANGE_LOOKUP = "Amfe"
RANGE_IDS = "AmfeIds"

COLUMNS = [
    Column("id_modo", "ID de modo", 14, styles.READONLY),
    Column("id_item", "ID de ítem", 14, styles.INPUT_TEXT),
    Column("item_desc", "Ítem", 40, styles.READONLY),
    Column("funcion", "Función", 46, styles.INPUT_TEXT),
    Column("falla_funcional", "Falla funcional", 40, styles.INPUT_TEXT),
    Column("modo_falla", "Modo de falla", 40, styles.INPUT_TEXT),
    Column("causa", "Causa", 40, styles.INPUT_TEXT),
    Column("efecto", "Efecto", 52, styles.INPUT_TEXT),
    Column("evidente", "¿Evidente al operador?", 14, styles.INPUT_TEXT, catalog="SI_NO"),
    Column("seguridad", "¿Afecta la seguridad?", 14, styles.INPUT_TEXT, catalog="SI_NO"),
    Column("ambiental", "¿Afecta el ambiente?", 14, styles.INPUT_TEXT, catalog="SI_NO"),
    Column("operacional", "¿Afecta la operación?", 14, styles.INPUT_TEXT, catalog="SI_NO"),
    Column("h_eje", "eje", 8, styles.READONLY),
    Column("h_evidencia", "evidencia", 8, styles.READONLY),
    Column("h_oculta", "oculta", 8, styles.READONLY),
    Column("categoria_consecuencia", "Categoría de consecuencia", 24, styles.READONLY),
    Column("codigo_modo_iso", "Modo ISO 14224", 14, styles.INPUT_TEXT),
    Column("notas", "Notas", 36, styles.INPUT_TEXT),
]

HIDDEN_KEYS = ("h_eje", "h_evidencia", "h_oculta")


class AmfeSheetBuilder:
    """Builds the functional failure analysis sheet."""

    spec = SPEC

    def build(self, worksheet: Worksheet, context: BuildContext) -> None:
        layout = SheetLayout(worksheet, first_column=2)
        layout.title(
            "Análisis de modos y efectos de falla",
            "Un registro por modo de falla. Las cuatro preguntas de consecuencia "
            "alimentan el árbol de decisión.",
        )
        header.build_echo(worksheet, layout)
        layout.note(
            "Responda las cuatro preguntas con S o N. La categoría de consecuencia "
            "se deriva sola y no se edita."
        )
        layout.blank()

        region = layout.table(COLUMNS, [], blank_rows=DATA_ROWS, autofilter=True, freeze=True)
        column_of = {
            column.key: region.first_column + offset for offset, column in enumerate(COLUMNS)
        }

        self._write_formulas(worksheet, region, column_of, context)
        self._hide_helpers(worksheet, column_of)
        self._validations(worksheet, region, column_of)
        self._define_names(context, region, column_of)

        worksheet.column_dimensions["A"].width = 3

    # -- formulas --------------------------------------------------------

    def _write_formulas(
        self, worksheet: Worksheet, region, column_of: dict[str, int], context: BuildContext
    ) -> None:
        letter = {key: get_column_letter(column) for key, column in column_of.items()}
        description_index = context.regions["particion.columna_descripcion"]

        for row in range(region.first_data_row, region.last_data_row + 1):
            mode = f"{letter['modo_falla']}{row}"
            item = f"{letter['id_item']}{row}"
            evident = f"{letter['evidente']}{row}"
            safety = f"{letter['seguridad']}{row}"
            environment = f"{letter['ambiental']}{row}"
            operation = f"{letter['operacional']}{row}"

            worksheet.cell(
                row=row, column=column_of["id_modo"]
            ).value = f'=IF({mode}="","",PARAM_PREFIJO_MODO&TEXT(ROW()-{region.header_row},"0000"))'
            worksheet.cell(row=row, column=column_of["item_desc"]).value = (
                f'=IF({item}="","",'
                f"IFERROR(VLOOKUP({item},RNG_PARTICION,{description_index},FALSE),"
                f'"ítem no encontrado"))'
            )

            worksheet.cell(row=row, column=column_of["h_eje"]).value = (
                f'=IF({safety}="S","SEG",'
                f'IF({environment}="S","AMB",'
                f'IF({operation}="S","OPE","NOOPE")))'
            )
            worksheet.cell(
                row=row, column=column_of["h_evidencia"]
            ).value = f'=IF({evident}="S","EVIDENTE","OCULTA")'
            # A hidden failure only splits two ways, so the OR lives here instead
            # of inside the category formula, which would then nest four deep.
            worksheet.cell(
                row=row, column=column_of["h_oculta"]
            ).value = f'=IF(OR({safety}="S",{environment}="S"),"OCULTA-SEG","OCULTA-NOSEG")'

            axis = f"{letter['h_eje']}{row}"
            evidence = f"{letter['h_evidencia']}{row}"
            hidden = f"{letter['h_oculta']}{row}"
            worksheet.cell(row=row, column=column_of["categoria_consecuencia"]).value = (
                f'=IF({evident}="","",IF({evidence}="OCULTA",{hidden},{evidence}&"-"&{axis}))'
            )

    @staticmethod
    def _hide_helpers(worksheet: Worksheet, column_of: dict[str, int]) -> None:
        for key in HIDDEN_KEYS:
            worksheet.column_dimensions[get_column_letter(column_of[key])].hidden = True

    # -- validations -----------------------------------------------------

    @staticmethod
    def _validations(worksheet: Worksheet, region, column_of: dict[str, int]) -> None:
        for column in COLUMNS:
            if column.catalog:
                add_list_validation(
                    worksheet,
                    names.catalog_name(column.catalog),
                    region.first_data_row,
                    region.last_data_row,
                    column_of[column.key],
                )

        item_column = get_column_letter(column_of["id_item"])
        add_custom_validation(
            worksheet,
            f"COUNTIF(RNG_PARTICIONIDS,{item_column}{region.first_data_row})>0",
            region.first_data_row,
            region.last_data_row,
            column_of["id_item"],
            error_title="Ítem inexistente",
            error_message="El ID de ítem debe existir en la hoja Particion.",
        )

    # -- names -----------------------------------------------------------

    def _define_names(self, context: BuildContext, region, column_of: dict[str, int]) -> None:
        names.define(
            context.workbook,
            names.range_name(RANGE_LOOKUP),
            names.range_ref(
                self.spec.title,
                region.first_column,
                region.first_data_row,
                region.last_column,
                region.last_data_row,
            ),
        )
        names.define(
            context.workbook,
            names.range_name(RANGE_IDS),
            names.range_ref(
                self.spec.title,
                column_of["id_modo"],
                region.first_data_row,
                column_of["id_modo"],
                region.last_data_row,
            ),
        )
        for key in (
            "funcion",
            "falla_funcional",
            "modo_falla",
            "causa",
            "efecto",
            "categoria_consecuencia",
            "evidente",
            "seguridad",
            "ambiental",
            "operacional",
        ):
            context.publish(f"amfe.columna_{key}", str(column_of[key] - region.first_column + 1))
