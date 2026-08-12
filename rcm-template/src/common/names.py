"""Defined names: the stable addresses the VBA modules and formulas rely on.

Nothing in the workbook may hard code a cell address. Every value a formula or
a macro needs is reachable through a defined name built here, so a layout change
never breaks a consumer.

Conventions:

  PARAM_<clave>   a single parameter value cell
  CAT_<catalogo>  a catalogue code range, used as a data validation source
  LIB_<nombre>    a library data range
  RNG_<nombre>    any other named range

Ranges are sized exactly to their seeded content and use absolute references, so
no volatile function is needed to resolve them.
"""

from __future__ import annotations

import re

from openpyxl.utils import get_column_letter, quote_sheetname
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.workbook.workbook import Workbook

PREFIX_PARAM = "PARAM_"
PREFIX_CATALOG = "CAT_"
PREFIX_LIBRARY = "LIB_"
PREFIX_RANGE = "RNG_"

# Excel allows letters, digits, underscore and period, and forbids a leading digit.
_INVALID_CHARS = re.compile(r"[^A-Za-z0-9_.]")


class NameError_(ValueError):
    """Raised when a defined name would be invalid or duplicated."""


def sanitize(fragment: str) -> str:
    """Turn an arbitrary key into a fragment usable inside a defined name."""
    cleaned = _INVALID_CHARS.sub("_", fragment.strip())
    if not cleaned:
        raise NameError_(f"Cannot build a defined name from '{fragment}'.")
    return cleaned.upper()


def cell_ref(sheet_title: str, column: int, row: int) -> str:
    """Absolute reference to a single cell, e.g. `Parametros!$D$12`."""
    return f"{quote_sheetname(sheet_title)}!${get_column_letter(column)}${row}"


def range_ref(
    sheet_title: str, first_column: int, first_row: int, last_column: int, last_row: int
) -> str:
    """Absolute reference to a rectangular range."""
    start = f"${get_column_letter(first_column)}${first_row}"
    end = f"${get_column_letter(last_column)}${last_row}"
    return f"{quote_sheetname(sheet_title)}!{start}:{end}"


def define(workbook: Workbook, name: str, reference: str) -> None:
    """Register a workbook scoped defined name, refusing duplicates."""
    if name in workbook.defined_names:
        raise NameError_(f"Defined name '{name}' already exists.")
    workbook.defined_names[name] = DefinedName(name, attr_text=reference)


def param_name(clave: str) -> str:
    return PREFIX_PARAM + sanitize(clave)


def catalog_name(catalogo: str) -> str:
    return PREFIX_CATALOG + sanitize(catalogo)


def library_name(nombre: str) -> str:
    return PREFIX_LIBRARY + sanitize(nombre)


def range_name(nombre: str) -> str:
    return PREFIX_RANGE + sanitize(nombre)


def formula_ref(name: str) -> str:
    """A defined name used inside a formula string."""
    return name
