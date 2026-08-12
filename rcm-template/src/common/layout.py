"""Layout primitives shared by the sheet builders.

Everything visual that more than one sheet needs lives here: titles, section
bands, header rows, column widths, freeze panes and list validations. Builders
describe *what* goes on the sheet; this module decides how it looks.
"""

from __future__ import annotations

from dataclasses import dataclass, field
from typing import Any

from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.worksheet import Worksheet

from src.common import styles

# Rows above the first table on every sheet: title, subtitle, blank.
TITLE_ROW = 2
SUBTITLE_ROW = 3


@dataclass(frozen=True)
class Column:
    """One column of a table: its heading, width and the style of its cells."""

    key: str
    header: str
    width: float = 18
    style: str = styles.INPUT_TEXT
    catalog: str | None = None
    note: str = ""


@dataclass(frozen=True)
class TableRegion:
    """Where a table ended up, so callers can build defined names over it."""

    header_row: int
    first_data_row: int
    last_data_row: int
    first_column: int
    last_column: int

    @property
    def row_count(self) -> int:
        return self.last_data_row - self.first_data_row + 1


@dataclass
class SheetLayout:
    """Cursor over a sheet, so builders never track row numbers by hand."""

    worksheet: Worksheet
    row: int = TITLE_ROW
    first_column: int = 2
    _validations: list[DataValidation] = field(default_factory=list)

    def title(self, text: str, subtitle: str = "") -> None:
        cell = self.worksheet.cell(row=TITLE_ROW, column=self.first_column, value=text)
        cell.style = styles.TITLE
        self.row = TITLE_ROW + 1
        if subtitle:
            note = self.worksheet.cell(row=SUBTITLE_ROW, column=self.first_column, value=subtitle)
            note.style = styles.SUBTITLE
            self.row = SUBTITLE_ROW + 1
        self.row += 1
        self.worksheet.sheet_view.showGridLines = False

    def section(self, text: str, width: int = 6) -> int:
        """Write a section band and return the row it occupies."""
        row = self.row
        for offset in range(width):
            cell = self.worksheet.cell(row=row, column=self.first_column + offset)
            cell.style = styles.SECTION
        self.worksheet.cell(row=row, column=self.first_column, value=text)
        self.worksheet.row_dimensions[row].height = 18
        self.row = row + 1
        return row

    def note(self, text: str) -> int:
        row = self.row
        cell = self.worksheet.cell(row=row, column=self.first_column, value=text)
        cell.style = styles.NOTE
        self.row = row + 1
        return row

    def blank(self, count: int = 1) -> None:
        self.row += count

    def table(
        self,
        columns: list[Column],
        rows: list[dict[str, Any]],
        *,
        blank_rows: int = 0,
        autofilter: bool = False,
        freeze: bool = False,
    ) -> TableRegion:
        """Write a header row followed by data rows, and return the region."""
        worksheet = self.worksheet
        header_row = self.row
        first_column = self.first_column
        last_column = first_column + len(columns) - 1

        for offset, column in enumerate(columns):
            cell = worksheet.cell(row=header_row, column=first_column + offset, value=column.header)
            cell.style = styles.HEADER
            if column.note:
                cell.comment = None  # comments are added by phase 9 tooling, not here
            worksheet.column_dimensions[
                get_column_letter(first_column + offset)
            ].width = column.width
        worksheet.row_dimensions[header_row].height = 30

        first_data_row = header_row + 1
        total_rows = len(rows) + blank_rows
        for index in range(total_rows):
            source = rows[index] if index < len(rows) else {}
            for offset, column in enumerate(columns):
                cell = worksheet.cell(row=first_data_row + index, column=first_column + offset)
                cell.style = column.style
                value = source.get(column.key)
                if value not in (None, ""):
                    cell.value = value

        last_data_row = first_data_row + max(total_rows, 1) - 1
        region = TableRegion(
            header_row=header_row,
            first_data_row=first_data_row,
            last_data_row=last_data_row,
            first_column=first_column,
            last_column=last_column,
        )

        if autofilter:
            worksheet.auto_filter.ref = (
                f"{get_column_letter(first_column)}{header_row}:"
                f"{get_column_letter(last_column)}{last_data_row}"
            )
        if freeze:
            worksheet.freeze_panes = worksheet.cell(row=first_data_row, column=first_column)

        self.row = last_data_row + 2
        return region

    def key_values(
        self,
        entries: list[tuple[str, Any, str]],
        *,
        label_width: float = 34,
        value_width: float = 28,
        note_width: float = 60,
        value_style: str = styles.INPUT_TEXT,
    ) -> TableRegion:
        """Write a label / value / note block and return its region.

        Each entry is `(label, value, note)`. Used by `Parametros` and by the
        identification header of the analysis sheets.
        """
        worksheet = self.worksheet
        first_row = self.row
        label_col = self.first_column
        value_col = label_col + 1
        note_col = label_col + 2

        worksheet.column_dimensions[get_column_letter(label_col)].width = label_width
        worksheet.column_dimensions[get_column_letter(value_col)].width = value_width
        worksheet.column_dimensions[get_column_letter(note_col)].width = note_width

        for index, (label, value, note) in enumerate(entries):
            row = first_row + index
            label_cell = worksheet.cell(row=row, column=label_col, value=label)
            label_cell.style = styles.LABEL
            value_cell = worksheet.cell(row=row, column=value_col)
            value_cell.style = value_style
            if value not in (None, ""):
                value_cell.value = value
            note_cell = worksheet.cell(row=row, column=note_col)
            note_cell.style = styles.NOTE
            if note:
                note_cell.value = note

        last_row = first_row + max(len(entries), 1) - 1
        self.row = last_row + 2
        return TableRegion(
            header_row=first_row - 1,
            first_data_row=first_row,
            last_data_row=last_row,
            first_column=label_col,
            last_column=note_col,
        )


def add_list_validation(
    worksheet: Worksheet,
    defined_name: str,
    first_row: int,
    last_row: int,
    column: int,
    *,
    allow_blank: bool = True,
    error_message: str = "El valor debe pertenecer a la lista.",
) -> DataValidation:
    """Restrict a column range to the codes of a defined name."""
    validation = DataValidation(
        type="list",
        formula1=defined_name,
        allow_blank=allow_blank,
        showDropDown=False,  # False means "show the dropdown arrow" in the file format
        showErrorMessage=True,
        errorTitle="Valor no permitido",
        error=error_message,
    )
    worksheet.add_data_validation(validation)
    letter = get_column_letter(column)
    validation.add(f"{letter}{first_row}:{letter}{last_row}")
    return validation


def add_range_validation(
    worksheet: Worksheet,
    first_row: int,
    last_row: int,
    column: int,
    *,
    minimum: float | str,
    maximum: float | str,
    decimal: bool = False,
    error_message: str = "",
) -> DataValidation:
    """Restrict a column range to a numeric interval.

    Bounds may be defined names rather than literals, which is how the criticality
    matrix keeps its scale in `Parametros` instead of on the sheet.
    """
    validation = DataValidation(
        type="decimal" if decimal else "whole",
        operator="between",
        formula1=str(minimum),
        formula2=str(maximum),
        allow_blank=True,
        showErrorMessage=True,
        errorTitle="Valor fuera de rango",
        error=error_message or f"El valor debe estar entre {minimum} y {maximum}.",
    )
    worksheet.add_data_validation(validation)
    letter = get_column_letter(column)
    validation.add(f"{letter}{first_row}:{letter}{last_row}")
    return validation


def add_custom_validation(
    worksheet: Worksheet,
    formula: str,
    first_row: int,
    last_row: int,
    column: int,
    *,
    error_title: str,
    error_message: str,
) -> DataValidation:
    """Accept a value only when `formula` evaluates true for it.

    Used where a list would be unusable -- a reference to a column of mostly
    empty rows, for instance -- but the value still has to exist somewhere.
    """
    validation = DataValidation(
        type="custom",
        formula1=formula,
        allow_blank=True,
        showErrorMessage=True,
        errorTitle=error_title,
        error=error_message,
    )
    worksheet.add_data_validation(validation)
    letter = get_column_letter(column)
    validation.add(f"{letter}{first_row}:{letter}{last_row}")
    return validation
