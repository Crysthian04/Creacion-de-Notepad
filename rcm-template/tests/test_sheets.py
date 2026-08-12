"""Structural tests of the built workbook.

These assert the contract every sheet owes the rest of the workbook: the sheets
exist, the defined names resolve, the validations point at real catalogues, and
no formula breaks the project's own rules.
"""

from __future__ import annotations

import pytest
from openpyxl import load_workbook

from src.common import names
from src.common.formulas import MAX_NESTING, distinct_formulas, scan_workbook
from src.common.workbook import build_workbook
from src.sheets import registry
from src.sheets.base import VERY_HIDDEN, VISIBLE

EXPECTED_SHEETS = (
    "INICIO",
    "Contexto Operacional",
    "Criticidad del Activo",
    "Particion",
    "AMFE",
    "Parametros",
    "Diccionario",
    "Lib_Taxonomia",
    "Lib_Particiones",
    "Lib_ModosFalla",
    "Lib_Tareas",
)


@pytest.fixture(scope="module")
def workbook():
    """One built workbook shared by the whole module; building it is the slow part."""
    from src.common.config import load_config

    built, _ = build_workbook(load_config())
    return built


@pytest.fixture(scope="module")
def formulas(workbook):
    return scan_workbook(workbook)


# --------------------------------------------------------------------------
# Sheets
# --------------------------------------------------------------------------


def test_every_expected_sheet_is_present(workbook):
    assert tuple(workbook.sheetnames) == EXPECTED_SHEETS


def test_sheet_keys_are_unique():
    keys = [builder.spec.key for builder in registry()]
    assert len(keys) == len(set(keys))


def test_libraries_are_very_hidden(workbook):
    for title in ("Diccionario", "Lib_Taxonomia", "Lib_Particiones", "Lib_ModosFalla"):
        assert workbook[title].sheet_state == VERY_HIDDEN


def test_working_sheets_are_visible(workbook):
    for title in ("INICIO", "Contexto Operacional", "Particion", "AMFE"):
        assert workbook[title].sheet_state == VISIBLE


# --------------------------------------------------------------------------
# Defined names
# --------------------------------------------------------------------------


def test_every_parameter_has_a_defined_name(workbook, config):
    from src.seedloader import load_seed

    seed = load_seed(config.project_root / "seed")
    for row in seed.parametros:
        assert names.param_name(row["clave"]) in workbook.defined_names


def test_every_catalogue_has_a_defined_name(workbook, config):
    from src.seedloader import load_seed

    seed = load_seed(config.project_root / "seed")
    for catalogue in seed.catalogos_por_nombre:
        assert names.catalog_name(catalogue) in workbook.defined_names


def test_catalogue_ranges_are_sized_to_their_content(workbook, config):
    from src.seedloader import load_seed

    seed = load_seed(config.project_root / "seed")
    for catalogue, entries in seed.catalogos_por_nombre.items():
        reference = workbook.defined_names[names.catalog_name(catalogue)].attr_text
        first, last = reference.split("!")[1].split(":")
        first_row = int(first.split("$")[2])
        last_row = int(last.split("$")[2])
        assert last_row - first_row + 1 == len(entries), catalogue


def test_defined_names_use_absolute_references(workbook):
    for name, defined in workbook.defined_names.items():
        assert "$" in defined.attr_text, name


# --------------------------------------------------------------------------
# Formulas
# --------------------------------------------------------------------------


def test_no_formula_references_an_unknown_name(workbook, formulas):
    known = set(workbook.defined_names)
    offenders = [
        (info.sheet, info.cell, sorted(info.names - known))
        for info in distinct_formulas(formulas)
        if info.names - known
    ]
    assert not offenders, f"Formulas reference undefined names: {offenders}"


def test_no_formula_nests_deeper_than_the_limit(formulas):
    offenders = [
        (info.sheet, info.cell, info.depth, info.formula)
        for info in distinct_formulas(formulas)
        if info.depth > MAX_NESTING
    ]
    assert not offenders, f"Formulas nested deeper than {MAX_NESTING}: {offenders}"


def test_no_formula_uses_a_volatile_function(formulas):
    # OFFSET, INDIRECT, TODAY, NOW and RAND force a full recalculation chain and
    # are banned by the performance constraint.
    volatile = ("OFFSET(", "INDIRECT(", "TODAY(", "NOW(", "RAND(")
    offenders = [
        (info.sheet, info.cell, info.formula)
        for info in distinct_formulas(formulas)
        if any(token in info.formula.upper() for token in volatile)
    ]
    assert not offenders, f"Volatile functions found: {offenders}"


def test_no_formula_references_a_whole_column(formulas):
    offenders = [
        (info.sheet, info.cell, info.formula)
        for info in distinct_formulas(formulas)
        if any(f"${letter}:${letter}" in info.formula for letter in "ABCDEFGHIJKLMNOPQRSTUVWXYZ")
    ]
    assert not offenders, f"Whole-column references found: {offenders}"


def test_the_workbook_actually_contains_formulas(formulas):
    assert len(formulas) > 100


# --------------------------------------------------------------------------
# Validations
# --------------------------------------------------------------------------


def test_list_validations_point_at_defined_names(workbook):
    known = set(workbook.defined_names)
    for worksheet in workbook.worksheets:
        for validation in worksheet.data_validations.dataValidation:
            if validation.type != "list":
                continue
            source = (validation.formula1 or "").lstrip("=")
            assert source in known, f"{worksheet.title}: unknown list source '{source}'"


def test_amfe_restricts_the_consequence_answers(workbook):
    worksheet = workbook["AMFE"]
    sources = {
        (validation.formula1 or "").lstrip("=")
        for validation in worksheet.data_validations.dataValidation
        if validation.type == "list"
    }
    assert names.catalog_name("SI_NO") in sources


def test_criticality_scale_comes_from_parameters(workbook):
    worksheet = workbook["Criticidad del Activo"]
    bounds = {
        (validation.formula1, validation.formula2)
        for validation in worksheet.data_validations.dataValidation
        if validation.type == "whole"
    }
    assert ("PARAM_ESCALA_MIN", "PARAM_ESCALA_MAX") in bounds


# --------------------------------------------------------------------------
# Round trip through a real file
# --------------------------------------------------------------------------


def test_workbook_survives_a_save_and_reload(tmp_path, config):
    from src.common.workbook import write_workbook

    built, _ = build_workbook(config)
    path = write_workbook(built, tmp_path / "roundtrip.xlsx")
    reloaded = load_workbook(path)

    assert tuple(reloaded.sheetnames) == EXPECTED_SHEETS
    assert set(reloaded.defined_names) == set(built.defined_names)
