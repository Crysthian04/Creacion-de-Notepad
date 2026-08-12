"""Phase 1 tests: the build pipeline itself.

These cover everything that does not need Excel. The COM stage is asserted only
through its skip path; running it for real requires Windows and is documented in
docs/entorno-build.md.
"""

from __future__ import annotations

import json
from pathlib import Path

import pytest
import yaml
from openpyxl import load_workbook

from build.build import (
    BuildError,
    build_workbook,
    file_sha256,
    main,
    parse_args,
    resolve_output_path,
    run_build,
)
from src.common.config import ConfigError, load_config


def _silent(*_args, **_kwargs) -> None:
    """Log sink so the build under test does not spam the test output."""


# --------------------------------------------------------------------------
# Configuration
# --------------------------------------------------------------------------


def test_config_loads_and_is_complete(config):
    assert config.version
    assert config.workbook.default_output.endswith(".xlsm")
    assert config.vba.components, "At least one VBA component must be declared."


def test_declared_vba_components_all_exist(config):
    paths = config.vba_component_paths()
    assert paths
    assert all(path.is_file() for path in paths)


def test_config_rejects_missing_component(tmp_path: Path, project_root: Path):
    raw = yaml.safe_load((project_root / "build" / "build.config.yaml").read_text("utf-8"))
    raw["vba"]["components"] = ["modDoesNotExist.bas"]
    broken = tmp_path / "build.config.yaml"
    broken.write_text(yaml.safe_dump(raw), encoding="utf-8")

    config = load_config(broken, project_root=project_root)
    with pytest.raises(ConfigError, match="does not exist"):
        config.vba_component_paths()


def test_config_rejects_missing_key(tmp_path: Path, project_root: Path):
    broken = tmp_path / "build.config.yaml"
    broken.write_text(yaml.safe_dump({"version": "0.0.1"}), encoding="utf-8")
    with pytest.raises(ConfigError, match="workbook"):
        load_config(broken, project_root=project_root)


def test_config_rejects_missing_file(tmp_path: Path):
    with pytest.raises(ConfigError, match="not found"):
        load_config(tmp_path / "nope.yaml")


# --------------------------------------------------------------------------
# Workbook assembly (stage 1)
# --------------------------------------------------------------------------


def test_workbook_has_registered_sheets(config):
    workbook, specs = build_workbook(config)
    assert specs, "The registry must produce at least one sheet."
    assert workbook.sheetnames == [spec.title for spec in specs]


def test_workbook_has_no_leftover_default_sheet(config):
    workbook, _ = build_workbook(config)
    assert "Sheet" not in workbook.sheetnames


def test_workbook_properties_carry_the_version(config):
    workbook, _ = build_workbook(config)
    assert workbook.properties.version == config.version
    assert workbook.properties.title == config.workbook.properties["title"]


def test_sheet_titles_fit_excel_limits(config):
    _, specs = build_workbook(config)
    for spec in specs:
        assert 1 <= len(spec.title) <= 31


# --------------------------------------------------------------------------
# Full build without the COM stage
# --------------------------------------------------------------------------


@pytest.fixture()
def built(config, tmp_path: Path):
    output = tmp_path / "RCM_Template_test.xlsm"
    result = run_build(config, output, use_com=False, log=_silent)
    return result, tmp_path


def test_build_without_com_produces_a_readable_xlsx(built):
    result, _ = built
    assert result.output_path.suffix == ".xlsx"
    assert result.output_path.is_file()

    workbook = load_workbook(result.output_path)
    assert workbook.sheetnames == result.sheet_titles


def test_build_without_com_reports_the_skip(built):
    result, _ = built
    assert result.com_stage == "omitida (--no-com)"
    assert result.imported_components == []
    assert result.signed is False


def test_manifest_is_written_next_to_the_output(built):
    result, tmp_path = built
    manifest_path = tmp_path / "manifest.json"
    assert manifest_path.is_file()

    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    assert manifest["version"] == result.version
    assert manifest["sheets"]["count"] == result.sheet_count
    assert manifest["output"]["macro_enabled"] is False
    assert manifest["output"]["sha256"] == file_sha256(result.output_path)
    assert manifest["built_at_utc"].endswith("+00:00")


def test_manifest_hashes_every_declared_vba_module(built, config):
    _, tmp_path = built
    manifest = json.loads((tmp_path / "manifest.json").read_text(encoding="utf-8"))

    hashed = {entry["file"]: entry["sha256"] for entry in manifest["vba_modules"]}
    assert set(hashed) == set(config.vba.components)
    for path in config.vba_component_paths():
        assert hashed[path.name] == file_sha256(path)


def test_build_is_repeatable(config, tmp_path: Path):
    first = run_build(config, tmp_path / "a.xlsm", use_com=False, log=_silent)
    second = run_build(config, tmp_path / "a.xlsm", use_com=False, log=_silent)
    assert first.sheet_titles == second.sheet_titles
    assert second.output_path.is_file()


def test_require_com_fails_when_excel_is_unavailable(config, tmp_path: Path):
    from build import preflight

    if preflight.check_environment().can_run_com:
        pytest.skip("This machine can run the COM stage.")

    with pytest.raises(BuildError, match="--require-com"):
        run_build(config, tmp_path / "a.xlsm", use_com=True, require_com=True, log=_silent)


def test_sign_without_com_stage_is_refused(config, tmp_path: Path):
    with pytest.raises(BuildError, match="firmar"):
        run_build(config, tmp_path / "a.xlsm", use_com=False, sign_project=True, log=_silent)


# --------------------------------------------------------------------------
# Preflight
# --------------------------------------------------------------------------


def test_preflight_blocks_the_com_stage_off_windows():
    from build import preflight

    report = preflight.check_environment()
    if preflight.is_windows():
        pytest.skip("Windows-specific behaviour is asserted on Windows only.")

    assert not report.can_run_com
    assert any(issue.code == "not-windows" for issue in report.blockers)


def test_preflight_report_names_the_remedy():
    from build import preflight

    report = preflight.check_environment()
    text = preflight.format_report(report)
    for issue in report.issues:
        assert issue.remedy in text


# --------------------------------------------------------------------------
# CLI
# --------------------------------------------------------------------------


def test_cli_defaults_to_the_configured_output(config):
    args = parse_args([])
    assert resolve_output_path(config, args.output) == config.default_output_path()


def test_cli_expands_the_version_placeholder(config):
    path = resolve_output_path(config, "dist/RCM_{version}.xlsm")
    assert path.name == f"RCM_{config.version}.xlsm"


def test_cli_returns_zero_without_com(tmp_path: Path, capsys):
    code = main(["--no-com", "--output", str(tmp_path / "cli.xlsm")])
    captured = capsys.readouterr()
    assert code == 0, captured.err
    assert "OK" in captured.out


def test_cli_reports_failure_without_traceback(tmp_path: Path, capsys):
    code = main(["--config", str(tmp_path / "missing.yaml")])
    captured = capsys.readouterr()
    assert code == 1
    assert "Build fallido" in captured.err
