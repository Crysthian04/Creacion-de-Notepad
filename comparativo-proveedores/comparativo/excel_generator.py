"""Genera una versión editable en Excel (.xlsx) del comparativo, con estilos
equivalentes (colores, bordes, negrilla) a los medidos en la plantilla PDF.

Nota: no se recibió un archivo .xlsx original como plantilla, solo el PDF de
referencia. Esta hoja recrea la misma estructura y estilo visual para que sea
editable, pero no proviene de copiar un libro de Excel existente.
"""

from __future__ import annotations

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from . import layout as L
from .modelos import Comparativo

GRIS_ENCABEZADO = "A6A6A6"
GRIS_FONDO = "F2F2F2"
BLANCO = "FFFFFF"

BORDE_FINO = Side(style="thin", color="000000")
BORDE = Border(left=BORDE_FINO, right=BORDE_FINO, top=BORDE_FINO, bottom=BORDE_FINO)


def _aplicar_borde(ws, celda_rango):
    for fila in ws[celda_rango]:
        for celda in fila:
            celda.border = BORDE


def generar_excel(comparativo: Comparativo, ruta_salida: str) -> None:
    n = len(comparativo.proveedores)

    wb = Workbook()
    ws = wb.active
    ws.title = "Selección Proveedores"
    ws.sheet_view.showGridLines = False

    # Columnas: A=Criterios, B=Peso, luego 2 columnas (Cot / %Ponderado) por proveedor
    ws.column_dimensions["A"].width = 42
    ws.column_dimensions["B"].width = 8
    col = 3
    for _ in range(n):
        ws.column_dimensions[get_column_letter(col)].width = 14
        ws.column_dimensions[get_column_letter(col + 1)].width = 14
        col += 2

    ultima_col_datos = 2 + n * 2
    # Se reserva 1 columna extra (más allá de los datos) para la caja de código,
    # de modo que el título y la caja de código nunca se queden sin espacio
    # aunque haya un solo proveedor.
    ultima_col = ultima_col_datos + 1

    def celda(fila, columna):
        return ws.cell(row=fila, column=columna)

    # ------------------------------------------------------------------
    # Logo + título + código
    # ------------------------------------------------------------------
    ws.row_dimensions[1].height = 45
    try:
        img = XLImage(L.LOGO_PATH)
        img.height = 55
        img.width = 108
        ws.add_image(img, "A1")
    except Exception:
        pass

    ws.merge_cells(start_row=1, start_column=3, end_row=1, end_column=ultima_col_datos)
    c = celda(1, 3)
    c.value = "SELECCIÓN DE PROVEEDORES LAC"
    c.font = Font(bold=True, size=14)
    c.alignment = Alignment(horizontal="center", vertical="center")

    ws.merge_cells(start_row=1, start_column=ultima_col, end_row=1, end_column=ultima_col)
    c = celda(1, ultima_col)
    c.value = "CÓDIGO: PSN-LAC-LAC-F-009 | VERSION: 3"
    c.font = Font(bold=True, size=8)
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.column_dimensions[get_column_letter(ultima_col)].width = 22

    # ------------------------------------------------------------------
    # Fecha / Bien-Servicio / Negociador
    # ------------------------------------------------------------------
    fila = 3
    for etiqueta, valor in [
        ("Fecha", comparativo.fecha),
        ("Bien/Servicio", comparativo.bien_o_servicio),
        ("Negociador/ Comprador Funcional", comparativo.negociador),
    ]:
        c = celda(fila, 1)
        c.value = etiqueta
        c.font = Font(bold=False)
        ws.merge_cells(start_row=fila, start_column=2, end_row=fila, end_column=4)
        cv = celda(fila, 2)
        cv.value = valor
        cv.alignment = Alignment(horizontal="center")
        cv.fill = PatternFill("solid", fgColor=BLANCO)
        fila += 1

    # ------------------------------------------------------------------
    # Encabezado de tabla
    # ------------------------------------------------------------------
    fila_header = fila + 1
    c = celda(fila_header, 1)
    c.value = "CRITERIOS"
    c.font = Font(bold=True)
    c.fill = PatternFill("solid", fgColor=GRIS_ENCABEZADO)
    c.alignment = Alignment(horizontal="center", vertical="center")

    c = celda(fila_header, 2)
    c.value = "Peso"
    c.font = Font(bold=True)
    c.fill = PatternFill("solid", fgColor=GRIS_ENCABEZADO)
    c.alignment = Alignment(horizontal="center", vertical="center")

    columna = 3
    for prov in comparativo.proveedores:
        ws.merge_cells(start_row=fila_header, start_column=columna,
                        end_row=fila_header, end_column=columna + 1)
        c = celda(fila_header, columna)
        c.value = prov.nombre
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor=GRIS_ENCABEZADO)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        celda(fila_header, columna + 1).fill = PatternFill("solid", fgColor=GRIS_ENCABEZADO)
        columna += 2

    fila_sub = fila_header + 1
    columna = 3
    for _ in comparativo.proveedores:
        celda(fila_sub, columna).value = "Cot."
        celda(fila_sub, columna + 1).value = "% Ponderado"
        for offset in (0, 1):
            celda(fila_sub, columna + offset).font = Font(size=9)
            celda(fila_sub, columna + offset).alignment = Alignment(horizontal="center")
        columna += 2

    # ------------------------------------------------------------------
    # Filas de criterios
    # ------------------------------------------------------------------
    valores_crudos = {
        "legalmente_constituido": lambda p: p.legalmente_constituido,
        "cumplimiento_especificaciones": lambda p: p.cumplimiento_especificaciones,
        "precio": lambda p: p.precio,
        "terminos_pago": lambda p: p.terminos_pago,
        "tiempo_respuesta": lambda p: p.tiempo_respuesta,
    }
    ponderados = {
        "legalmente_constituido": lambda p: p.legalmente_constituido_ponderado,
        "cumplimiento_especificaciones": lambda p: p.cumplimiento_especificaciones_ponderado,
        "precio": lambda p: p.precio_ponderado,
        "terminos_pago": lambda p: p.terminos_pago_ponderado,
        "tiempo_respuesta": lambda p: p.tiempo_respuesta_ponderado,
    }

    fila_criterio_inicio = fila_sub + 1
    fila = fila_criterio_inicio
    for clave, etiqueta, peso in L.CRITERIOS:
        celda(fila, 1).value = etiqueta
        c = celda(fila, 2)
        c.value = peso / 100
        c.number_format = "0%"
        c.font = Font(bold=True)
        c.alignment = Alignment(horizontal="center")

        columna = 3
        for prov in comparativo.proveedores:
            cv = celda(fila, columna)
            valor = valores_crudos[clave](prov)
            if clave == "precio":
                cv.value = valor
                cv.number_format = '"$"#,##0.00'
            else:
                cv.value = valor
            cv.alignment = Alignment(horizontal="center")

            cp = celda(fila, columna + 1)
            cp.value = ponderados[clave](prov) / 100
            cp.number_format = "0%"
            cp.font = Font(bold=True)
            cp.alignment = Alignment(horizontal="center")
            columna += 2
        fila += 1

    # ------------------------------------------------------------------
    # Fila REQ + totales
    # ------------------------------------------------------------------
    fila_req = fila
    ws.row_dimensions[fila_req].height = 30
    ws.merge_cells(start_row=fila_req, start_column=1, end_row=fila_req, end_column=1)
    c = celda(fila_req, 1)
    c.value = f"{comparativo.req_numero} {comparativo.req_descripcion}"
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    c = celda(fila_req, 2)
    c.value = L.PESO_TOTAL / 100
    c.number_format = "0%"
    c.font = Font(bold=True)
    c.alignment = Alignment(horizontal="center", vertical="center")

    columna = 3
    for prov in comparativo.proveedores:
        c = celda(fila_req, columna + 1)
        c.value = prov.total_ponderado() / 100
        c.number_format = "0%"
        c.font = Font(bold=True)
        c.alignment = Alignment(horizontal="center", vertical="center")
        columna += 2

    # ------------------------------------------------------------------
    # Proveedor sugerido / autorizado / observaciones
    # ------------------------------------------------------------------
    fila = fila_req + 2
    for etiqueta, valor in [
        ("Proveedor Sugerido", comparativo.proveedor_sugerido),
        ("Proveedor Autorizado", comparativo.proveedor_autorizado),
    ]:
        celda(fila, 1).value = etiqueta
        ws.merge_cells(start_row=fila, start_column=2, end_row=fila, end_column=4)
        celda(fila, 2).value = valor
        fila += 1

    fila += 1
    celda(fila, 1).value = "Observaciones Relevantes:"
    celda(fila, 1).font = Font(bold=True)
    fila += 1
    ws.merge_cells(start_row=fila, start_column=1, end_row=fila + 3, end_column=ultima_col)
    c = celda(fila, 1)
    c.value = comparativo.observaciones
    c.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)

    # ------------------------------------------------------------------
    # Bordes
    # ------------------------------------------------------------------
    ultima_col_letra = get_column_letter(ultima_col)
    _aplicar_borde(ws, f"A{fila_header}:{ultima_col_letra}{fila_req}")

    wb.save(ruta_salida)
